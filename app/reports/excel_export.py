from __future__ import annotations

import json
import shutil
import warnings
from datetime import datetime
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker
from openpyxl.styles import Font
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.cell import coordinate_from_string

from app.assinaturas import resolver_assinatura
from app.config import ROOT
from app.models.receipt import ReceiptStatus
from app.models.trip import Trip
from app.paths import bundled_asset, user_asset

PASTA_MODELO_RDV = user_asset("rdv")
ARQUIVO_MAPEAMENTO = PASTA_MODELO_RDV / "mapeamento_celulas.json"

# Fallback se o JSON não existir ou estiver incompleto
_CABECALHO_PADRAO: dict[str, dict[str, Any]] = {
    "contratada": {"celula": "C1", "modo": "texto"},
    "inicio_contratual": {"celula": "G2", "modo": "data", "formato": "DD/MM/YYYY"},
    "termino_contratual": {"celula": "G3", "modo": "data", "formato": "DD/MM/YYYY"},
    "numero_rdv": {"celula": "G4", "modo": "rotulo", "prefixo": "RDV-"},
    "nome_funcionario": {"celula": "C5", "modo": "rotulo", "prefixo": "Colaborador: "},
    "contratante": {"celula": "F5", "modo": "rotulo", "prefixo": "Cliente: "},
    "area": {"celula": "C6", "modo": "rotulo", "prefixo": "Area: "},
    "empreendimento": {"celula": "F6", "modo": "rotulo", "prefixo": "Projeto: "},
    "numero_os": {"celula": "C7", "modo": "rotulo", "prefixo": "Nº O S.: ", "padrao_vazio": "N/D"},
    "local": {
        "celula": "F7",
        "modo": "composicao",
        "prefixo": "Local: ",
        "partes": [
            {"campos": ["endereco"]},
            {"campos": ["cidade", "estado"], "separador": "/"},
        ],
        "separador_partes": ", ",
    },
    "adiantamento": {"celula": "I10", "modo": "moeda", "formato": "#,##0.00"},
}

_CATEGORIAS_PADRAO: dict[str, int] = {
    "Hospedagem": 14,
    "Estacionamento / garagem": 27,
    "Café da manhã": 16,
    "Frigobar": 18,
    "Lavanderia": 19,
    "Alimentação": 20,
    "Táxi / aplicativo": 23,
    "Locação de veículo": 25,
    "Combustível": 26,
    "Pedágio": 28,
    "Bagagem": 32,
    "Passagem": 31,
    "Outros": 34,
}


def _parse_iso_date(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        return datetime.strptime(value.strip()[:10], "%Y-%m-%d")
    except ValueError:
        return None


@lru_cache(maxsize=1)
def load_mapeamento_rdv() -> dict[str, Any]:
    """Carrega o mapeamento do Excel (cópia do usuário, senão o empacotado)."""
    candidatos = (
        ARQUIVO_MAPEAMENTO,
        bundled_asset("rdv", "mapeamento_celulas.json"),
    )
    for path in candidatos:
        if path.is_file():
            with open(path, encoding="utf-8") as f:
                data = json.load(f) or {}
            if isinstance(data, dict):
                return data
    return {}


def reload_mapeamento_rdv() -> dict[str, Any]:
    """Recarrega o JSON (útil após editar o arquivo com o app aberto)."""
    load_mapeamento_rdv.cache_clear()
    return load_mapeamento_rdv()


def _template_path(mapeamento: dict[str, Any] | None = None) -> Path:
    mapeamento = mapeamento if mapeamento is not None else load_mapeamento_rdv()
    nome = str(mapeamento.get("arquivo_modelo") or "RDV-PADRÃO.xlsx").strip() or "RDV-PADRÃO.xlsx"
    candidatos = [
        PASTA_MODELO_RDV / nome,
        bundled_asset("rdv", nome),
        ROOT / nome,  # legado: modelo na raiz do projeto
    ]
    for path in candidatos:
        if path.is_file():
            return path
    return candidatos[0]


def _valor_composicao(trip: Trip, cfg: dict[str, Any]) -> str:
    """Monta texto a partir de vários campos do projeto (ex.: endereço + cidade/estado)."""
    sep_partes = str(cfg.get("separador_partes") if cfg.get("separador_partes") is not None else ", ")
    partes_cfg = cfg.get("partes")
    montadas: list[str] = []

    if isinstance(partes_cfg, list) and partes_cfg:
        for parte in partes_cfg:
            if not isinstance(parte, dict):
                continue
            campos = parte.get("campos") or []
            sep = str(parte.get("separador") if parte.get("separador") is not None else " ")
            pedacos = [
                str(getattr(trip, str(c), None) or "").strip()
                for c in campos
                if str(c).strip()
            ]
            pedacos = [p for p in pedacos if p]
            if pedacos:
                montadas.append(sep.join(pedacos))
    else:
        # Fallback: lista simples de campos
        campos = cfg.get("campos") or ["endereco", "cidade", "estado"]
        sep = str(cfg.get("separador") if cfg.get("separador") is not None else ", ")
        pedacos = [
            str(getattr(trip, str(c), None) or "").strip()
            for c in campos
            if str(c).strip()
        ]
        montadas = [p for p in pedacos if p]
        return sep.join(montadas)

    return sep_partes.join(montadas)


def _fill_rdv_header(ws, trip: Trip, mapeamento: dict[str, Any]) -> None:
    cabecalho = mapeamento.get("cabecalho") or _CABECALHO_PADRAO
    for campo, cfg in cabecalho.items():
        if campo.startswith("_") or not isinstance(cfg, dict):
            continue
        celula = str(cfg.get("celula") or "").strip()
        if not celula:
            continue
        modo = str(cfg.get("modo") or "texto").strip().lower()
        valor_bruto = getattr(trip, campo, None) if hasattr(trip, campo) else None

        if modo == "data":
            dt = _parse_iso_date(str(valor_bruto) if valor_bruto else None)
            ws[celula] = dt if dt else None
            if dt:
                ws[celula].number_format = str(cfg.get("formato") or "DD/MM/YYYY")
            continue

        if modo == "moeda":
            if valor_bruto is None:
                continue
            ws[celula] = float(valor_bruto)
            ws[celula].number_format = str(cfg.get("formato") or "#,##0.00")
            continue

        if modo == "composicao":
            prefixo = str(cfg.get("prefixo") or "")
            texto = _valor_composicao(trip, cfg)
            padrao = cfg.get("padrao_vazio")
            if not texto and padrao is not None:
                texto = str(padrao)
            ws[celula] = f"{prefixo}{texto}".rstrip()
            continue

        if modo == "rotulo":
            prefixo = str(cfg.get("prefixo") or "")
            texto = str(valor_bruto or "").strip()
            padrao = cfg.get("padrao_vazio")
            if not texto and padrao is not None:
                texto = str(padrao)
            # Evita duplicar prefixo se o usuário já digitou (ex.: RDV-12345)
            if prefixo and texto.upper().startswith(prefixo.upper()):
                texto = texto[len(prefixo) :].lstrip()
            if not texto and not padrao:
                continue
            ws[celula] = f"{prefixo}{texto}".rstrip()
            continue

        # modo texto
        texto = str(valor_bruto or "").strip()
        if not texto:
            continue
        ws[celula] = texto


def _clear_sample_marks(ws, mapeamento: dict[str, Any]) -> None:
    """Remove marcas de exemplo do modelo (ex.: 'x' de taxa de adm)."""
    cfg = mapeamento.get("limpar_marcas_exemplo") or {}
    coluna = str(cfg.get("coluna") or "B").strip() or "B"
    linhas = cfg.get("linhas") or [14, 25, 31]
    marca = cfg.get("valor_marca", "x")
    for row in linhas:
        ref = f"{coluna}{int(row)}"
        if ws[ref].value == marca:
            ws[ref] = None


def _fill_rdv_totais(ws, trip: Trip, mapeamento: dict[str, Any]) -> None:
    from app.categorias_usuario import mapa_rotulo_custom, mesclar_linhas_excel

    totais_cfg = mapeamento.get("totais_categoria") or {}
    coluna = str(totais_cfg.get("coluna") or "G").strip() or "G"
    formato = str(totais_cfg.get("formato") or "#,##0.00")
    linhas_map = mesclar_linhas_excel(totais_cfg.get("linhas") or _CATEGORIAS_PADRAO)
    categoria_padrao = str(totais_cfg.get("categoria_padrao") or "Outros")
    if categoria_padrao not in linhas_map:
        categoria_padrao = next(iter(linhas_map), "Outros")

    rotulos_custom = mapa_rotulo_custom()
    totais = trip.totais_por_categoria()

    for row in set(int(v) for v in linhas_map.values()):
        ws[f"{coluna}{row}"] = None

    for cat, total in totais.items():
        row = linhas_map.get(cat)
        if row is None:
            row = linhas_map.get(categoria_padrao)
        if row is None:
            continue
        row = int(row)
        ref = f"{coluna}{row}"
        atual = ws[ref].value
        try:
            atual_f = float(atual or 0)
        except (TypeError, ValueError):
            atual_f = 0.0
        ws[ref] = atual_f + float(total)
        ws[ref].number_format = formato

        # Região Diversas: escreve o nome da categoria custom ao lado do valor
        if cat in rotulos_custom:
            ws[rotulos_custom[cat]] = cat


def _fill_rdv_observacoes(wb, trip: Trip, mapeamento: dict[str, Any]) -> None:
    cfg = mapeamento.get("observacoes") or {}
    aba = str(cfg.get("aba") or mapeamento.get("aba_verso") or "verso")
    celula = str(cfg.get("celula") or "A2").strip() or "A2"
    if aba not in wb.sheetnames:
        return
    obs = (trip.observacoes or "").strip()
    if not obs:
        return
    wb[aba][celula] = obs


def _col_width_px(ws, col_letter: str) -> float:
    dim = ws.column_dimensions.get(col_letter)
    width = dim.width if dim is not None and dim.width is not None else None
    if width is None:
        width = ws.sheet_format.defaultColWidth or 8.43
    # Conversão usual do Excel (Calibri 11): px ≈ largura * 7 + 5
    return float(width) * 7.0 + 5.0


def _row_height_px(ws, row: int) -> float:
    dim = ws.row_dimensions.get(row)
    height = dim.height if dim is not None and dim.height is not None else None
    if height is None:
        height = ws.sheet_format.defaultRowHeight or 15.0
    # Altura em pontos → pixels (96 dpi)
    return float(height) * (96.0 / 72.0)


def _intervalo_tamanho_px(ws, inicio: str, fim: str) -> tuple[int, int]:
    c1, r1 = coordinate_from_string(inicio.upper())
    c2, r2 = coordinate_from_string(fim.upper())
    i1 = column_index_from_string(c1)
    i2 = column_index_from_string(c2)
    if i2 < i1:
        i1, i2 = i2, i1
    if r2 < r1:
        r1, r2 = r2, r1
    largura = sum(_col_width_px(ws, get_column_letter(i)) for i in range(i1, i2 + 1))
    altura = sum(_row_height_px(ws, r) for r in range(int(r1), int(r2) + 1))
    return max(1, int(round(largura))), max(1, int(round(altura)))


def _aparar_assinatura(pil_img):
    """Remove margem transparente / fundo branco para a tinta preencher melhor a área."""
    import numpy as np
    from PIL import Image as PILImage

    img = pil_img.convert("RGBA")
    arr = np.asarray(img)
    if arr.size == 0:
        return img
    a = arr[:, :, 3]
    r, g, b = arr[:, :, 0], arr[:, :, 1], arr[:, :, 2]
    mask = (a >= 20) & ~((r >= 248) & (g >= 248) & (b >= 248))
    ys, xs = np.where(mask)
    if len(xs) == 0:
        return img
    pad = 3
    h, w = arr.shape[:2]
    min_x = max(0, int(xs.min()) - pad)
    min_y = max(0, int(ys.min()) - pad)
    max_x = min(w, int(xs.max()) + 1 + pad)
    max_y = min(h, int(ys.max()) + 1 + pad)
    return PILImage.fromarray(arr[min_y:max_y, min_x:max_x])


def _anchor_marker(celula: str, *, col_off_px: int = 0, row_off_px: int = 0) -> AnchorMarker:
    from openpyxl.utils.units import pixels_to_EMU

    col_letters, row = coordinate_from_string(celula.upper())
    col = column_index_from_string(col_letters) - 1
    row_idx = int(row) - 1
    return AnchorMarker(
        col=col,
        colOff=pixels_to_EMU(max(0, col_off_px)),
        row=row_idx,
        rowOff=pixels_to_EMU(max(0, row_off_px)),
    )


def _fill_rdv_assinatura(wb, trip: Trip, mapeamento: dict[str, Any]) -> None:
    """Insere a assinatura em D60:E61 sem distorcer, alinhada à linha inferior."""
    import io

    from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor
    from openpyxl.drawing.xdr import XDRPositiveSize2D
    from openpyxl.utils.units import pixels_to_EMU
    from PIL import Image as PILImage

    cfg = mapeamento.get("assinatura") or {}
    aba = str(cfg.get("aba") or mapeamento.get("aba_frente") or "frente")
    if aba not in wb.sheetnames:
        return
    inicio = str(cfg.get("celula_inicio") or "D60").strip().upper() or "D60"
    fim = str(cfg.get("celula_fim") or "E61").strip().upper() or "E61"

    path = resolver_assinatura(trip.nome_funcionario, trip.assinatura_arquivo)
    if not path:
        return

    ws = wb[aba]

    # Garante altura mínima nas linhas da assinatura (só no arquivo gerado)
    c1, r1 = coordinate_from_string(inicio)
    c2, r2 = coordinate_from_string(fim)
    if r2 < r1:
        r1, r2 = r2, r1
    altura_minima_pt = float(cfg.get("altura_minima_linha_pt") or 28)
    for r in range(int(r1), int(r2) + 1):
        atual = ws.row_dimensions[r].height
        if atual is None or float(atual) < altura_minima_pt:
            ws.row_dimensions[r].height = altura_minima_pt

    box_w, box_h = _intervalo_tamanho_px(ws, inicio, fim)
    # Margens internas (esquerda/direita e folga acima da linha inferior)
    margem_x = 10
    margem_topo = 4
    margem_base = 4
    util_w = max(1, box_w - 2 * margem_x)
    util_h = max(1, box_h - margem_topo - margem_base)

    try:
        with PILImage.open(path) as raw:
            aparada = _aparar_assinatura(raw)
            # Redimensiona já na imagem (melhor nitidez no Excel)
            escala = min(util_w / aparada.width, util_h / aparada.height)
            dest_w = max(1, int(round(aparada.width * escala)))
            dest_h = max(1, int(round(aparada.height * escala)))
            if dest_w != aparada.width or dest_h != aparada.height:
                aparada = aparada.resize((dest_w, dest_h), PILImage.Resampling.LANCZOS)
            buf = io.BytesIO()
            aparada.save(buf, format="PNG")
            buf.seek(0)
    except Exception:
        return

    try:
        img = XLImage(buf)
    except Exception:
        return
    img.width = dest_w
    img.height = dest_h

    # Centraliza na horizontal; alinha embaixo (sobre a linha de assinatura)
    off_x = margem_x + max(0, (util_w - dest_w) // 2)
    off_y = max(0, box_h - margem_base - dest_h)
    img.anchor = OneCellAnchor(
        _from=_anchor_marker(inicio, col_off_px=off_x, row_off_px=off_y),
        ext=XDRPositiveSize2D(pixels_to_EMU(dest_w), pixels_to_EMU(dest_h)),
    )
    ws.add_image(img)


def _fill_rdv_data_geracao(wb, mapeamento: dict[str, Any]) -> None:
    """Preenche a(s) célula(s) de Data com a data atual (ex.: C60)."""
    cfg = mapeamento.get("data_geracao") or {}
    aba = str(cfg.get("aba") or mapeamento.get("aba_frente") or "frente")
    if aba not in wb.sheetnames:
        return
    celulas = cfg.get("celulas") or ["C60"]
    if isinstance(celulas, str):
        celulas = [celulas]
    formato = str(cfg.get("formato") or "DD/MM/YYYY")
    hoje = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    ws = wb[aba]
    for ref in celulas:
        celula = str(ref or "").strip().upper()
        if not celula:
            continue
        ws[celula] = hoje
        ws[celula].number_format = formato


def _add_lancamentos_sheet(wb, trip: Trip) -> None:
    if "Lançamentos" in wb.sheetnames:
        del wb["Lançamentos"]
    ws = wb.create_sheet("Lançamentos")
    headers = ["Data", "Categoria", "Estabelecimento", "Valor", "Moeda", "Status", "Arquivo", "Confiança", "Fonte"]
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(1, col, h)
        cell.font = Font(bold=True)

    ativos = [r for r in trip.receipts if r.status != ReceiptStatus.EXCLUIDA]
    ativos.sort(key=lambda r: (r.data_ordenacao(), r.estabelecimento or ""))

    for i, r in enumerate(ativos, start=2):
        ws.cell(i, 1, r.data or "")
        ws.cell(i, 2, r.categoria)
        ws.cell(i, 3, r.estabelecimento)
        valor_cell = ws.cell(i, 4, r.valor if r.valor is not None else "")
        if r.valor is not None:
            valor_cell.number_format = '"R$" #,##0.00'
        ws.cell(i, 5, r.moeda)
        ws.cell(i, 6, r.status.value)
        ws.cell(i, 7, Path(r.arquivo).name)
        ws.cell(i, 8, round(r.confianca, 2))
        ws.cell(i, 9, r.fonte_extracao)

    widths = [12, 24, 32, 14, 8, 12, 28, 10, 12]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def _export_excel_simples(trip: Trip, output_path: Path) -> Path:
    """Fallback se o modelo RDV não existir."""
    wb = Workbook()
    ws_resumo = wb.active
    ws_resumo.title = "Resumo"
    ws_resumo["A1"] = "Relatório de Despesas de Viagem"
    ws_resumo["A1"].font = Font(bold=True, size=14)
    rows_meta = [
        ("Natureza do serviço", trip.natureza_servico),
        ("Empreendimento", trip.empreendimento),
        ("Contratante", trip.contratante),
        ("Endereço", trip.endereco),
        ("Cidade", trip.cidade),
        ("Estado", trip.estado),
        ("Nome funcionário", trip.nome_funcionario),
        ("Contratada", trip.contratada),
        ("Área", trip.area),
        ("Nº OS", trip.numero_os),
        ("Nº RDV", trip.numero_rdv),
        ("Período", f"{trip.inicio_contratual or '—'} a {trip.termino_contratual or '—'}"),
        ("Centro de custo", trip.centro_custo),
        ("Adiantamento", trip.adiantamento if trip.adiantamento is not None else 0),
        ("Total geral", trip.total_geral()),
        ("A reembolsar", trip.valor_a_reembolsar()),
    ]
    for i, (label, value) in enumerate(rows_meta, start=3):
        ws_resumo.cell(i, 1, label).font = Font(bold=True)
        cell = ws_resumo.cell(i, 2, value)
        if label in ("Total geral", "Adiantamento", "A reembolsar") and isinstance(value, (int, float)):
            cell.number_format = '"R$" #,##0.00'
    _add_lancamentos_sheet(wb, trip)
    wb.save(output_path)
    return output_path


def export_excel(trip: Trip, output_path: Path) -> Path:
    """Gera o Excel RDV a partir do modelo em assets/rdv/."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    mapeamento = reload_mapeamento_rdv()
    template = _template_path(mapeamento)
    if not template.exists():
        return _export_excel_simples(trip, output_path)

    # Cópia do modelo preserva formatação/fórmulas
    shutil.copy2(template, output_path)

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        wb = load_workbook(output_path)

    aba_frente = str(mapeamento.get("aba_frente") or "frente")
    if aba_frente not in wb.sheetnames:
        return _export_excel_simples(trip, output_path)

    ws = wb[aba_frente]
    _fill_rdv_header(ws, trip, mapeamento)
    _clear_sample_marks(ws, mapeamento)
    _fill_rdv_totais(ws, trip, mapeamento)
    _fill_rdv_observacoes(wb, trip, mapeamento)
    _fill_rdv_data_geracao(wb, mapeamento)
    _fill_rdv_assinatura(wb, trip, mapeamento)
    _add_lancamentos_sheet(wb, trip)

    wb.save(output_path)
    return output_path
